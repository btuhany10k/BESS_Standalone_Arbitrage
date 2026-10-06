"""
BESS 0.5C - Kapsamlı Excel Rapor Üreticisi (.xlsx)
8.760 Saatlik Detay, 365 Günlük Özet, 12 Aylık Kırılım, KPI ve Yıllık Kıyaslama sayfalarını içerir.
"""

import io
from typing import Dict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
from .optimizer_05c import BESSConfig05C


def generate_05c_excel_report(
    year: int,
    config: BESSConfig05C,
    hourly_df: pd.DataFrame,
    daily_df: pd.DataFrame,
    kpis: Dict,
    monthly_df: pd.DataFrame,
    comp_df: pd.DataFrame,
) -> bytes:
    """5 sayfalı 0.5C BESS Fizibilite ve Arbitraj Çalışma Kitabını bellek üzerinde oluşturur."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Varsayılan boş sayfayı sil

    # Renk Paleti (Senior Industrial Dark/Navy Theme)
    NAVY_HEADER = "0F172A"
    CYAN_ACCENT = "0284C7"
    GRAY_SUB = "334155"
    BORDER_LIGHT = "CBD5E1"
    
    font_title = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Segoe UI", size=10)
    font_bold = Font(name="Segoe UI", size=10, bold=True)
    
    fill_header = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    fill_accent = PatternFill(start_color=CYAN_ACCENT, end_color=CYAN_ACCENT, fill_type="solid")
    fill_sub = PatternFill(start_color=GRAY_SUB, end_color=GRAY_SUB, fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color=BORDER_LIGHT),
        right=Side(style="thin", color=BORDER_LIGHT),
        top=Side(style="thin", color=BORDER_LIGHT),
        bottom=Side(style="thin", color=BORDER_LIGHT)
    )

    # -------------------------------------------------------------
    # SAYFA 1: ÖZET & PARAMETRELER
    # -------------------------------------------------------------
    ws_summary = wb.create_sheet(title="Özet & Parametreler")
    ws_summary.views.sheetView[0].showGridLines = True

    ws_summary.merge_cells("A1:D1")
    title_cell = ws_summary["A1"]
    title_cell.value = f"BESS 0.5C (2 Saatlik Depolama) Arbitraj Fizibilite Raporu - {year}"
    title_cell.font = font_title
    title_cell.fill = fill_header
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_summary.row_dimensions[1].height = 35

    summary_rows = [
        ("Sistem Konfigürasyonu", ""),
        ("Analiz Yılı", year),
        ("Nominal Güç (MW)", f"{config.power_mw:.1f} MW"),
        ("C-Rate & Süre", "0.5C (2 Saatlik Depolama)"),
        ("Toplam Enerji Kapasitesi (MWh)", f"{config.capacity_mwh:.1f} MWh"),
        ("Döngü Verimliliği (RTE %)", f"%{config.rte * 100:.0f}"),
        ("Güne Başlangıç SoC (%)", f"%{config.soc_start_pct:.0f}"),
        ("Gün Sonu Hedef SoC (%)", f"%{config.soc_end_pct:.0f}"),
        ("Hücre Yıpranma Maliyeti ($/MWh)", f"${config.degradation_cost:.2f}"),
        ("Operasyon Stratejisi", "Sürekli Senaryo Modülasyonu" if "lp" in str(config.strategy) else ("Günde 2 Döngü (Çift Blok)" if config.strategy == "2_cycle" else "Günde 1 Döngü (2 Saat Blok)")),
        ("", ""),
        ("Yıllık Fizibilite ve Finansal KPI Sonuçları", ""),
        ("Net Arbitraj Kârı ($)", f"${kpis['net_profit']:,.2f}"),
        ("Brüt Arbitraj Kârı ($)", f"${kpis['gross_profit']:,.2f}"),
        ("Deşarj Satış Geliri ($)", f"${kpis['total_revenue']:,.2f}"),
        ("Şarj Elektrik Maliyeti ($)", f"${kpis['total_cost']:,.2f}"),
        ("Toplam Yıpranma Maliyeti ($)", f"${kpis['total_degradation_cost']:,.2f}"),
        ("Yıllık Eşdeğer Tam Döngü (EFC)", f"{kpis['total_cycles']:,.1f} EFC"),
        ("Döngü Başına Net Kâr ($/Döngü)", f"${kpis['profit_per_cycle']:,.2f}"),
        ("Aktif Operasyon Gün Sayısı", f"{kpis['active_days']} Gün"),
        ("Pas Geçilen / Bekleme Gün Sayısı", f"{kpis['passed_days']} Gün"),
        ("Ortalama Deşarj Fiyatı ($/MWh)", f"${kpis['avg_discharge_price']:.2f}"),
        ("Ortalama Şarj Fiyatı ($/MWh)", f"${kpis['avg_charge_price']:.2f}"),
        ("Gerçekleşen Fiyat Makası ($/MWh)", f"${kpis['realized_spread']:.2f}"),
    ]

    for idx, (label, val) in enumerate(summary_rows, start=3):
        c1 = ws_summary.cell(row=idx, column=1, value=label)
        c2 = ws_summary.cell(row=idx, column=2, value=val)
        if val == "":
            c1.font = Font(name="Segoe UI", size=11, bold=True, color="0284C7")
        else:
            c1.font = font_bold
            c2.font = font_data
            c1.border = thin_border
            c2.border = thin_border
        ws_summary.row_dimensions[idx].height = 20

    ws_summary.column_dimensions["A"].width = 38
    ws_summary.column_dimensions["B"].width = 32

    # -------------------------------------------------------------
    # SAYFA 2: 8760 SAATLİK DETAY
    # -------------------------------------------------------------
    ws_hourly = wb.create_sheet(title="8760 Saatlik Detay")
    ws_hourly.views.sheetView[0].showGridLines = True

    h_cols = [
        ("Tarih", 12), ("Saat", 8), ("PTF ($/MWh)", 14), ("Şarj Gücü (MW)", 15),
        ("Deşarj Gücü (MW)", 16), ("SoC (MWh)", 13), ("SoC (%)", 10),
        ("Şarj Maliyeti ($)", 16), ("Deşarj Geliri ($)", 16),
        ("Yıpranma ($)", 14), ("Net Kâr ($)", 14), ("İşlem Durumu", 26)
    ]
    for col_idx, (col_name, col_width) in enumerate(h_cols, start=1):
        cell = ws_hourly.cell(row=1, column=col_idx, value=col_name)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws_hourly.column_dimensions[get_column_letter(col_idx)].width = col_width
    ws_hourly.row_dimensions[1].height = 25

    # Saatlik verileri yaz (ilk 2000 satır veya tamamı - hızlı yanıt için 8760 satır)
    h_data = hourly_df[[
        "date", "saat_str", "ptf_usd", "p_ch_mw", "p_dis_mw", "soc_mwh", "soc_pct",
        "charge_cost", "discharge_revenue", "degradation_cost", "net_profit", "action_label"
    ]].values

    for r_idx, row_val in enumerate(h_data, start=2):
        for c_idx, val in enumerate(row_val, start=1):
            cell = ws_hourly.cell(row=r_idx, column=c_idx)
            if c_idx == 1:
                cell.value = str(val)[:10]
            elif isinstance(val, (int, float)):
                cell.value = round(val, 2)
                if c_idx in [3, 8, 9, 10, 11]:
                    cell.number_format = '$#,##0.00'
                elif c_idx == 7:
                    cell.number_format = '0.0"%"'
            else:
                cell.value = str(val)
            cell.font = font_data
            cell.border = thin_border

    # -------------------------------------------------------------
    # SAYFA 3: 365 GÜNLÜK ÖZET
    # -------------------------------------------------------------
    ws_daily = wb.create_sheet(title="365 Günlük Özet")
    ws_daily.views.sheetView[0].showGridLines = True

    d_cols = [
        ("Tarih", 12), ("PTF Spread ($)", 14), ("Deşarj Geliri ($)", 16), ("Şarj Maliyeti ($)", 16),
        ("Brüt Kâr ($)", 14), ("Yıpranma ($)", 14), ("Net Kâr ($)", 14), ("Döngü (EFC)", 12),
        ("Durum", 12)
    ]
    for col_idx, (col_name, col_width) in enumerate(d_cols, start=1):
        cell = ws_daily.cell(row=1, column=col_idx, value=col_name)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws_daily.column_dimensions[get_column_letter(col_idx)].width = col_width
    ws_daily.row_dimensions[1].height = 25

    d_data = daily_df[[
        "date", "ptf_spread", "discharge_revenue", "charge_cost", "gross_profit",
        "degradation_cost", "net_profit", "cycles", "is_passed"
    ]].values

    for r_idx, row_val in enumerate(d_data, start=2):
        for c_idx, val in enumerate(row_val, start=1):
            cell = ws_daily.cell(row=r_idx, column=c_idx)
            if c_idx == 1:
                cell.value = str(val)[:10]
            elif c_idx == 9:
                cell.value = "Pas Geçildi" if val else "Aktif"
            elif isinstance(val, (int, float)):
                cell.value = round(val, 2)
                if c_idx in [2, 3, 4, 5, 6, 7]:
                    cell.number_format = '$#,##0.00'
            cell.font = font_data
            cell.border = thin_border

    # -------------------------------------------------------------
    # SAYFA 4: AYLIK KIRILIM
    # -------------------------------------------------------------
    ws_monthly = wb.create_sheet(title="Aylık Kırılım")
    ws_monthly.views.sheetView[0].showGridLines = True

    m_cols = [
        ("Ay", 12), ("Deşarj Geliri ($)", 16), ("Şarj Maliyeti ($)", 16), ("Brüt Kâr ($)", 14),
        ("Yıpranma ($)", 14), ("Net Kâr ($)", 14), ("Toplam Döngü", 13), ("Aktif Gün", 11),
        ("Pas Gün", 11), ("$/Cycle", 13), ("Ort. Spread ($)", 15), ("Ort. PTF ($)", 14)
    ]
    for col_idx, (col_name, col_width) in enumerate(m_cols, start=1):
        cell = ws_monthly.cell(row=1, column=col_idx, value=col_name)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws_monthly.column_dimensions[get_column_letter(col_idx)].width = col_width
    ws_monthly.row_dimensions[1].height = 25

    m_data = monthly_df[[
        "month_name", "discharge_revenue", "charge_cost", "gross_profit", "degradation_cost",
        "net_profit", "cycles", "active_days", "passed_days", "profit_per_cycle", "avg_spread", "avg_ptf"
    ]].values

    for r_idx, row_val in enumerate(m_data, start=2):
        for c_idx, val in enumerate(row_val, start=1):
            cell = ws_monthly.cell(row=r_idx, column=c_idx)
            if isinstance(val, (int, float)):
                cell.value = round(val, 2)
                if c_idx in [2, 3, 4, 5, 6, 10, 11, 12]:
                    cell.number_format = '$#,##0.00'
            else:
                cell.value = str(val)
            cell.font = font_data
            cell.border = thin_border

    # -------------------------------------------------------------
    # SAYFA 5: 2024 vs 2025 KIYASLAMA
    # -------------------------------------------------------------
    ws_comp = wb.create_sheet(title="2024 vs 2025 Kıyaslama")
    ws_comp.views.sheetView[0].showGridLines = True

    comp_cols = list(comp_df.columns)
    for col_idx, col_name in enumerate(comp_cols, start=1):
        cell = ws_comp.cell(row=1, column=col_idx, value=col_name)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws_comp.column_dimensions[get_column_letter(col_idx)].width = 18
    ws_comp.row_dimensions[1].height = 25

    for r_idx, row_val in enumerate(comp_df.values, start=2):
        for c_idx, val in enumerate(row_val, start=1):
            cell = ws_comp.cell(row=r_idx, column=c_idx)
            if isinstance(val, (int, float)):
                cell.value = round(val, 2)
                if "$" in comp_cols[c_idx - 1]:
                    cell.number_format = '$#,##0.00'
            else:
                cell.value = str(val)
            cell.font = font_data
            cell.border = thin_border

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
